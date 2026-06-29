import base64
import io
from datetime import date, datetime

from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# Number formats
_FORMATS = {
    'text':       '@',
    'integer':    '#,##0',
    'float':      '#,##0.00',
    'monetary':   '#,##0.00',
    'date':       'YYYY-MM-DD',
    'datetime':   'YYYY-MM-DD HH:MM',
    'percentage': '0.0%',
    'boolean':    '@',
}


class ExcelExportWizard(models.TransientModel):
    _name = 'excel.export.wizard'
    _description = 'Advanced Excel Export Wizard'

    template_id = fields.Many2one(
        'excel.export.template',
        string='Template',
        required=True,
    )
    active_model = fields.Char(string='Model', readonly=True)
    active_ids_str = fields.Char(string='Selected IDs', readonly=True)
    record_count = fields.Integer(
        string='Records to Export',
        compute='_compute_record_count',
    )
    file_data = fields.Binary(string='Excel File', readonly=True)
    file_name = fields.Char(string='File Name', readonly=True)
    state = fields.Selection(
        [('choose', 'Choose'), ('done', 'Done')],
        default='choose',
    )

    @api.depends('active_ids_str')
    def _compute_record_count(self):
        for wiz in self:
            try:
                ids = [int(i) for i in (wiz.active_ids_str or '').split(',') if i.strip()]
                wiz.record_count = len(ids)
            except Exception:
                wiz.record_count = 0

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        ctx = self.env.context
        active_model = ctx.get('active_model', '')
        active_ids   = ctx.get('active_ids', [])

        res['active_model']   = active_model
        res['active_ids_str'] = ','.join(str(i) for i in active_ids)

        # Pre-select template if one matches the model
        if active_model:
            template = self.env['excel.export.template'].search([
                ('model_name', '=', active_model),
                ('active', '=', True),
            ], limit=1)
            if template:
                res['template_id'] = template.id

        return res

    # ------------------------------------------------------------------
    # Generate Excel
    # ------------------------------------------------------------------

    def action_open_wizard(self):
        """Return the window action to open this wizard form."""
        return {
            'type': 'ir.actions.act_window',
            'name': _('Export to Excel'),
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_export(self):
        self.ensure_one()
        if not HAS_OPENPYXL:
            raise UserError(_(
                'The openpyxl Python library is required for Excel export. '
                'Ask your system administrator to install it: '
                'pip install openpyxl'
            ))
        if not self.template_id.line_ids:
            raise UserError(_('The selected template has no columns defined.'))

        ids = [int(i) for i in self.active_ids_str.split(',') if i.strip()]
        if not ids:
            raise UserError(_('No records selected.'))

        records = self.env[self.active_model].browse(ids)
        template = self.template_id

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (template.sheet_title or template.name)[:31]

        # -- Header styles -----------------------------------------------
        h_fill = PatternFill('solid', fgColor=template.header_bg_color or '1F4E79')
        h_font = Font(
            bold=True,
            color=template.header_font_color or 'FFFFFF',
            name='Calibri',
            size=11,
        )
        h_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        alt_fill = PatternFill('solid', fgColor=template.alt_row_color or 'DDEEFF') \
            if template.alternating_rows else None
        thin = Side(style='thin', color='CCCCCC')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # -- Company header row (row 1) -----------------------------------
        row_offset = 0
        if template.include_company_header:
            company = self.env.company
            title = template.sheet_title or template.name
            header_text = f'{company.name}  —  {title}'
            ws.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=len(template.line_ids),
            )
            cell = ws.cell(1, 1, header_text)
            cell.font = Font(bold=True, size=13, name='Calibri',
                             color=template.header_bg_color or '1F4E79')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 24
            row_offset = 1

        # -- Column headers (row_offset + 1) ------------------------------
        header_row = row_offset + 1
        for col_idx, line in enumerate(template.line_ids, start=1):
            cell = ws.cell(header_row, col_idx,
                           line.header or line.field_name)
            cell.fill    = h_fill
            cell.font    = h_font
            cell.alignment = h_align
            cell.border  = border

        ws.row_dimensions[header_row].height = 22

        # -- Data rows ----------------------------------------------------
        totals = {}  # col_idx -> running sum

        for row_idx, record in enumerate(records, start=header_row + 1):
            use_alt = template.alternating_rows and (row_idx - header_row) % 2 == 0
            for col_idx, line in enumerate(template.line_ids, start=1):
                raw = self._get_field_value(record, line.field_name)
                value, fmt_str = self._format_value(raw, line.fmt)

                cell = ws.cell(row_idx, col_idx, value)
                cell.border = border
                if use_alt and alt_fill:
                    cell.fill = alt_fill
                if line.bold:
                    cell.font = Font(bold=True, name='Calibri')
                else:
                    cell.font = Font(name='Calibri')
                if fmt_str:
                    cell.number_format = fmt_str

                # Accumulate totals for numeric columns
                if line.total and isinstance(value, (int, float)):
                    totals[col_idx] = totals.get(col_idx, 0.0) + value

        # -- Totals row ---------------------------------------------------
        if totals:
            totals_row = header_row + len(records) + 1
            for col_idx, line in enumerate(template.line_ids, start=1):
                if col_idx in totals:
                    cell = ws.cell(totals_row, col_idx, totals[col_idx])
                    cell.font   = Font(bold=True, name='Calibri')
                    cell.fill   = PatternFill('solid', fgColor='D9E1F2')
                    cell.border = border
                    fmt_str = _FORMATS.get(line.fmt, '@')
                    if fmt_str != '@':
                        cell.number_format = fmt_str
                elif col_idx == 1:
                    cell = ws.cell(totals_row, col_idx, 'TOTAL')
                    cell.font = Font(bold=True, name='Calibri')
                    cell.fill = PatternFill('solid', fgColor='D9E1F2')
                    cell.border = border

        # -- Column widths ------------------------------------------------
        for col_idx, line in enumerate(template.line_ids, start=1):
            col_letter = get_column_letter(col_idx)
            if line.width and line.width > 0:
                ws.column_dimensions[col_letter].width = line.width
            else:
                # Auto-fit: measure max content length in column
                max_len = len(line.header or line.field_name) + 2
                for r in range(header_row + 1,
                               header_row + len(records) + 1):
                    v = ws.cell(r, col_idx).value
                    if v is not None:
                        max_len = max(max_len, len(str(v)) + 2)
                ws.column_dimensions[col_letter].width = min(max_len, 60)

        # -- Freeze header ------------------------------------------------
        if template.freeze_header:
            ws.freeze_panes = ws.cell(header_row + 1, 1)

        # -- Save to binary -----------------------------------------------
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = '%s.xlsx' % (
            (template.sheet_title or template.name).replace(' ', '_')
        )
        self.write({
            'file_data': base64.b64encode(buf.read()),
            'file_name': filename,
            'state':     'done',
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id':    self.id,
            'target':    'new',
            'context':   self.env.context,
        }

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _get_field_value(record, field_path):
        """Resolve dotted field paths like 'partner_id.name'."""
        try:
            parts = field_path.split('.')
            value = record
            for part in parts:
                if isinstance(value, models.BaseModel):
                    value = value[part]
                else:
                    return value
            # Collapse recordsets to display name
            if isinstance(value, models.BaseModel):
                if not value:
                    return ''
                return ', '.join(value.mapped('display_name'))
            return value
        except Exception:
            return ''

    @staticmethod
    def _format_value(raw, fmt):
        """Return (excel_value, number_format_string)."""
        fmt_str = _FORMATS.get(fmt, '@')

        if raw is False or raw is None:
            return '', None

        if fmt == 'boolean':
            return 'Yes' if raw else 'No', None

        if fmt == 'date':
            if isinstance(raw, datetime):
                return raw.date(), fmt_str
            if isinstance(raw, date):
                return raw, fmt_str
            return str(raw), None

        if fmt == 'datetime':
            if isinstance(raw, (date, datetime)):
                return raw, fmt_str
            return str(raw), None

        if fmt in ('integer', 'float', 'monetary', 'percentage'):
            try:
                return float(raw), fmt_str
            except (TypeError, ValueError):
                return str(raw), None

        return str(raw) if raw is not None else '', None
